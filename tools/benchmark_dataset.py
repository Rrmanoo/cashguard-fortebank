"""Private development benchmark. Labels are used only after inference, never in prompts."""
import argparse, asyncio, hashlib, json, os, sys, time
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))

def save(p,x):
 p=Path(p);p.parent.mkdir(parents=True,exist_ok=True);tmp=p.with_suffix('.tmp');tmp.write_text(json.dumps(x,ensure_ascii=False,indent=2));tmp.replace(p)
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
# Evaluation-only mapping from the owner's primary annotation, not exhaustive multilabel truth.

def references(path,type_map=None):
 from openpyxl import load_workbook
 book=load_workbook(path,data_only=True);sheet=book['labels'] if 'labels'in book.sheetnames else book.active;rows=[]
 for name,t,label,human_type,description,*_ in sheet.iter_rows(min_row=2,values_only=True):
  if not name:continue
  filename=str(name);filename=filename if filename.lower().endswith('.mov') else filename+'.mov';number=int(filename.rsplit('-',1)[1].split('.')[0]);positive=str(label).strip().lower() in ['нарушение','есть нарушение','violation','1','true','да']
  # Sheet uses Russian positive/negative words; reject unfamiliar labels instead of guessing.
  if str(label).strip().lower() not in ['нарушение','есть нарушение','violation','1','true','да','нет нарушения','без нарушения','no_violation','0','false','нет']:raise ValueError('Unknown label '+str(label))
  rows.append({'filename':filename,'expected':positive,'reference_timestamp_sec':t,'human_type':human_type,'human_description':description,'expected_types':(type_map or {}).get(str(number),[]) if positive else []})
 book.close();return rows

def metrics(rows):
 tp=tn=fp=fn=errors=hits=tags=neg_tags=0
 for r in rows:
  if r['status']!='done':errors+=1;continue
  p=r['verdict']=='violation';y=r['expected'];tp+=p and y;tn+=not p and not y;fp+=p and not y;fn+=not p and y
  hits+=len(set(r['expected_types'])&set(r['predicted_types']));tags+=len(r['predicted_types']);neg_tags+=len(r['predicted_types']) if not y else 0
 pos=sum(r['expected'] for r in rows);neg=len(rows)-pos;ref=sum(len(r['expected_types']) for r in rows)
 return {'n':len(rows),'positives':pos,'negatives':neg,'completed':len(rows)-errors,'failed':errors,'tp':tp,'tn':tn,'fp':fp,'fn_completed':fn,'positive_errors':sum(r['expected'] and r['status']!='done' for r in rows),'accuracy_all':(tp+tn)/len(rows),'precision':tp/(tp+fp) if tp+fp else 0,'recall_all':tp/pos if pos else None,'recall_completed':tp/(tp+fn) if tp+fn else None,'f1_all':2*tp/(2*tp+fp+pos-tp) if 2*tp+fp+pos-tp else 0,'specificity_all':tn/neg if neg else None,'balanced_accuracy_all':(tp/pos+tn/neg)/2 if pos and neg else None,'primary_reference_tags':ref,'primary_type_hits':hits,'primary_type_recall_all':hits/ref if ref else None,'predicted_type_tags':tags,'false_type_tags_on_negative_videos':neg_tags,'extra_types_on_positive_videos_are_unverified':True,'completed_under_300_sec':sum(r['status']=='done' and r['elapsed_sec']<300 for r in rows)}

async def run(a):
 from app.release.engine import Engine,analyze,VERSION
 from app.config import ROOT,MODEL_NAME,FRAME_MAX_EDGE
 from app.video import probe
 out=Path(a.output);out.mkdir(parents=True,exist_ok=True);rows=references(a.labels,json.loads(Path(a.type_map).read_text()) if a.type_map else None);videos=Path(a.videos);groups={}
 for r in rows:
  p=videos/r['filename'];r['sha256']=await asyncio.to_thread(sha,p);r['duration_sec']=(await asyncio.to_thread(probe,p))['duration'];groups.setdefault(r['sha256'],[]).append(r)
 source_hashes={str(p.relative_to(ROOT)):sha(p) for p in sorted(ROOT.rglob('*.py')) if any(part in ['app','rules'] for part in p.relative_to(ROOT).parts) and '__pycache__'not in p.parts}
 source_hashes.update({str(p.relative_to(ROOT)):sha(p) for p in (ROOT/'rules').glob('*.json')})
 save(out/'protocol.json',{'version':VERSION,'model_alias':MODEL_NAME,'frame_max_edge':FRAME_MAX_EDGE,'model_concurrency':4,'video_concurrency':a.concurrency,'deadline_sec':a.deadline,'fresh_unique_inferences':len(groups),'file_records':len(rows),'deduplication':'Exact SHA-256 bytes only; no previous model outputs reused','label_sha256':sha(a.labels),'source_hashes':source_hashes,'development_set':True,'labels_in_prompts':False,'thinking':False,'metric_policy':'Errors are separate; recall_all includes all positive records in denominator. Primary types are not exhaustive multilabel truth. Concurrent timings are not isolated SLA tests.'})
 save(out/'reference-labels.json',rows);done={};gate=asyncio.Semaphore(a.concurrency);start=time.monotonic()
 async with Engine() as engine:
  async def one(h,refs):
   async with gate:
    begin=time.monotonic();trace=out/'traces'/h[:16];print('START',len(done)+1,'/',len(groups),refs[0]['filename'],flush=True)
    try:
     result=await asyncio.wait_for(analyze(videos/refs[0]['filename'],engine=engine,trace=trace),timeout=a.deadline)
     record={'status':'done','verdict':result['verdict'],'predicted_types':[v['type_id'] for v in result['violations']],'result_path':str(trace/'result.json'),'elapsed_sec':round(time.monotonic()-begin,3)}
    except Exception as exc:
     record={'status':'failed','error':type(exc).__name__+': '+str(exc)[:250],'elapsed_sec':round(time.monotonic()-begin,3),'predicted_types':[]};save(trace/'failure.json',record)
    done[h]=record;expanded=[{**r,**done[r['sha256']]} for r in rows if r['sha256'] in done];save(out/'progress.json',{'unique_done':len(done),'unique_total':len(groups),'wall_seconds':time.monotonic()-start,'records':expanded,'metrics_so_far':metrics(expanded)});print('DONE',len(done),'/',len(groups),refs[0]['filename'],record['status'],round(record['elapsed_sec'],1),record.get('predicted_types',[]),flush=True)
  await asyncio.gather(*(one(h,refs) for h,refs in groups.items()))
 records=[{**r,**done[r['sha256']]} for r in rows];unique=[next(r for r in records if r['sha256']==h) for h in groups];times=sorted(r['elapsed_sec'] for r in unique);total={'version':VERSION,'wall_seconds':round(time.monotonic()-start,3),'sum_unique_job_seconds':round(sum(times),3),'video_seconds_file_weighted':sum(r['duration_sec'] for r in rows),'video_seconds_unique':sum(r['duration_sec'] for r in unique),'metrics_42':metrics(records),'metrics_40_unique':metrics(unique),'unique_job_seconds':{'min':times[0],'median':times[len(times)//2],'p95':times[min(len(times)-1,int(.95*len(times)))],'max':times[-1]},'duplicates':[[r['filename'] for r in refs] for refs in groups.values() if len(refs)>1],'records':records};save(out/'summary.json',total);print('FINISHED',json.dumps({k:v for k,v in total.items() if k!='records'},ensure_ascii=False),flush=True)

if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--videos',required=True);p.add_argument('--labels',required=True);p.add_argument('--output',required=True);p.add_argument('--type-map',help='Private evaluation-only JSON: numeric filename suffix -> annotated type IDs');p.add_argument('--concurrency',type=int,default=2);p.add_argument('--deadline',type=int,default=270);a=p.parse_args();asyncio.run(run(a))
